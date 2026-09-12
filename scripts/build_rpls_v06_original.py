import pandas as pd, numpy as np, shutil, zipfile, hashlib
from pathlib import Path
from scipy.stats import pearsonr, spearmanr
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
ROOT=Path('/mnt/data/FutureCareer_v0.6_longitudinal')
if ROOT.exists(): shutil.rmtree(ROOT)
for d in ['data','results','docs','scripts']: (ROOT/d).mkdir(parents=True,exist_ok=True)
emp=pd.read_csv('/mnt/data/employment_soc.csv'); post=pd.read_csv('/mnt/data/postings_by_occupation.csv'); sal=pd.read_csv('/mnt/data/salaries_soc.csv'); ha=pd.read_csv('/mnt/data/hiring_and_attrition_by_occupation.csv')
master=pd.read_csv('/mnt/data/FutureCareer_Research_Analysis_V3.1/Analysis_SOC6_Master_V3_1.csv')
for d in [emp,post,sal,ha]: d['month']=pd.to_datetime(d['month'])
for c in ['salary_nsa','salary_sa']: sal[c]=pd.to_numeric(sal[c].astype(str).str.replace('$','',regex=False).str.replace(',','',regex=False),errors='coerce')
post=post[post.soc2d_code.between(11,53)].copy(); sal=sal[sal.soc2d_code.between(11,53)].copy()
panel=emp.merge(post,on=['month','soc2d_code'],how='outer',suffixes=('_emp','_post')).merge(sal,on=['month','soc2d_code'],how='outer').merge(ha,on=['month','soc2d_code'],how='outer',suffixes=('','_ha'))
namecols=[c for c in panel.columns if 'soc2d_name' in c]; panel['soc2d_name']=panel[namecols].bfill(axis=1).iloc[:,0]; panel=panel.drop(columns=[c for c in namecols if c!='soc2d_name']).sort_values(['soc2d_code','month'])
panel.to_csv(ROOT/'data/RPLS_SOC2_Monthly_Panel_2021_2026.csv',index=False,encoding='utf-8-sig')
def trend(g,col):
 q=g[['month',col]].dropna().sort_values('month'); n=len(q)
 if n<12:return {'n_months':n,'start':np.nan,'end':np.nan,'growth_total_pct':np.nan,'annualized_log_trend_pct':np.nan}
 y=q[col].astype(float).values;t=(q.month-q.month.min()).dt.days.values/365.25;ok=y>0;s=np.polyfit(t[ok],np.log(y[ok]),1)[0] if ok.sum()>=12 else np.nan
 return {'n_months':n,'start':y[0],'end':y[-1],'growth_total_pct':(y[-1]/y[0]-1)*100 if y[0] else np.nan,'annualized_log_trend_pct':(np.exp(s)-1)*100 if pd.notna(s) else np.nan}
summ=[]
for soc,g in panel.groupby('soc2d_code'):
 r={'soc2d_code':int(soc),'soc2d_name':g.soc2d_name.dropna().iloc[0]}
 for key,col in [('employment','employment_sa'),('postings','active_postings_sa'),('salary','salary_sa')]: r.update({f'{key}_{k}':v for k,v in trend(g,col).items()})
 for col in ['rl_hiring_rate','rl_attrition_rate']:
  q=g[col].dropna();r[f'{col}_mean']=q.mean();r[f'{col}_latest']=q.iloc[-1] if len(q) else np.nan
 r['net_hiring_rate_mean']=r['rl_hiring_rate_mean']-r['rl_attrition_rate_mean'];summ.append(r)
trends=pd.DataFrame(summ);trends.to_csv(ROOT/'data/RPLS_SOC2_Trend_Summary.csv',index=False,encoding='utf-8-sig')
m=master.copy();m['soc2d_code']=pd.to_numeric(m.soc6.astype(str).str[:2],errors='coerce'); vars=['ilo_exposure_mean','observed_exposure','HumanMoat_seed_equal_weight','Physicality','Judgment','Interaction','Responsibility','Context'];rows=[]
for soc,g in m.groupby('soc2d_code'):
 r={'soc2d_code':int(soc),'soc6_count':g.soc6.nunique(),'employment2025_sum':pd.to_numeric(g['Employment, 2025'],errors='coerce').sum()};w=pd.to_numeric(g['Employment, 2025'],errors='coerce').fillna(0).clip(lower=0)
 for v in vars:
  x=pd.to_numeric(g[v],errors='coerce');ok=x.notna()&(w>0);r[v+'_weighted']=np.average(x[ok],weights=w[ok]) if ok.sum() and w[ok].sum()>0 else x.mean();r[v+'_unweighted']=x.mean()
 rows.append(r)
long=pd.DataFrame(rows).merge(trends,on='soc2d_code',how='inner');long.to_csv(ROOT/'data/SOC2_AI_Human_Labor_Longitudinal_Master.csv',index=False,encoding='utf-8-sig')
preds=['ilo_exposure_mean_weighted','observed_exposure_weighted','HumanMoat_seed_equal_weight_weighted']; outs=['employment_annualized_log_trend_pct','postings_annualized_log_trend_pct','salary_annualized_log_trend_pct','rl_hiring_rate_mean','rl_attrition_rate_mean','net_hiring_rate_mean'];res=[]
for x in preds:
 for y in outs:
  q=long[[x,y]].dropna()
  if len(q)>=10:
   pr,pp=pearsonr(q[x],q[y]);sr,sp=spearmanr(q[x],q[y]);res.append([x,y,len(q),pr,pp,sr,sp])
corr=pd.DataFrame(res,columns=['predictor','outcome','n','pearson_r','pearson_p','spearman_rho','spearman_p']);corr.to_csv(ROOT/'results/SOC2_Longitudinal_Correlations.csv',index=False)
coverage=pd.DataFrame([['Employment',emp.month.min().strftime('%Y-%m'),emp.month.max().strftime('%Y-%m'),emp.soc2d_code.nunique(),len(emp)],['Job openings',post.month.min().strftime('%Y-%m'),post.month.max().strftime('%Y-%m'),post.soc2d_code.nunique(),len(post)],['Salaries',sal.month.min().strftime('%Y-%m'),sal.month.max().strftime('%Y-%m'),sal.soc2d_code.nunique(),len(sal)],['Hiring & attrition',ha.month.min().strftime('%Y-%m'),ha.month.max().strftime('%Y-%m'),ha.soc2d_code.nunique(),len(ha)]],columns=['dataset','start','end','soc2_count','rows']);coverage.to_csv(ROOT/'results/RPLS_Coverage.csv',index=False)
wb=Workbook();wb.remove(wb.active)
for name,df in [('Coverage',coverage),('SOC2_Master',long),('Monthly_Panel',panel),('Correlations',corr)]:
 ws=wb.create_sheet(name)
 for j,c in enumerate(df.columns,1):cell=ws.cell(1,j,str(c));cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D');cell.alignment=Alignment(wrap_text=True)
 for i,row in enumerate(df.itertuples(index=False,name=None),2):
  for j,v in enumerate(row,1):ws.cell(i,j,None if pd.isna(v) else (int(v) if isinstance(v,np.integer) else float(v) if isinstance(v,np.floating) else v))
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for j,c in enumerate(df.columns,1):ws.column_dimensions[get_column_letter(j)].width=min(max(12,len(str(c))+2),32)
wb.save(ROOT/'FutureCareer_RPLS_Longitudinal_V0.6.xlsx')
(ROOT/'docs/RPLS_DATA_NOTE.md').write_text('''# RPLS Longitudinal Data Note\n\nFour public occupation timeseries were added. Employment and hiring/attrition cover 2021-01–2026-08; postings and salary cover 2022-01–2026-08. The grain is **SOC 2-digit**.\n\nFutureCareer therefore uses:\n\n```text\nSOC6 main analysis\n+ SOC2 longitudinal external validation\n```\n\nSeasonally adjusted series are primary. Annualized trends use a log-linear monthly slope:\n\n$$Trend_o=100(\\exp(\\hat\\beta_o)-1)$$\n\nSOC6 AI/Human measures are aggregated to SOC2 with 2025 employment weights:\n\n$$X_{SOC2}=\\frac{\\sum_j Employment_jX_j}{\\sum_j Employment_j}$$\n\nWith only about 22 SOC2 groups, correlations are exploratory, not causal. SOC2 values must never be presented as SOC6-specific historical facts.\n''',encoding='utf-8')
(ROOT/'docs/PROCESSING_STEPS.md').write_text('''# Processing Steps\n\n1. Read four raw CSVs unchanged.\n2. Parse month.\n3. Convert salary strings to dollars.\n4. Use seasonally adjusted fields.\n5. Remove Unknown/Other from postings/salary.\n6. Outer-join on month × SOC2.\n7. Compute annualized log trends for employment, postings and salary.\n8. Compute mean hiring, attrition and net hiring.\n9. Convert SOC6 research records to SOC2.\n10. Aggregate AI/Human Structure using BLS 2025 employment weights.\n11. Join SOC2 research measures to RPLS trends.\n12. Run Pearson/Spearman exploratory validation.\n13. Export panel, summaries, master, results and review workbook.\n\nNo missing longitudinal value is imputed.\n''',encoding='utf-8')
shutil.copy2('/mnt/data/build_v06.py',ROOT/'scripts/build_v06.py')
files=[]
for p in sorted(ROOT.rglob('*')):
 if p.is_file():files.append([str(p.relative_to(ROOT)),p.stat().st_size,hashlib.sha256(p.read_bytes()).hexdigest()])
pd.DataFrame(files,columns=['path','bytes','sha256']).to_csv(ROOT/'MANIFEST.csv',index=False)
with zipfile.ZipFile('/mnt/data/FutureCareer_RPLS_Longitudinal_V0.6.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in ROOT.rglob('*'):
  if p.is_file():z.write(p,p.relative_to(ROOT.parent))
print(coverage.to_string(index=False));print('\nSOC2 master',long.shape);print('\n',corr.to_string(index=False))
